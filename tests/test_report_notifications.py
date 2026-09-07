"""Stored-data report integration: PostgreSQL ledger + real renderers + fake mail."""

from concurrent.futures import ThreadPoolExecutor
from datetime import UTC, datetime, timedelta
from io import BytesIO
from types import SimpleNamespace

import pytest
from app.core.context import CurrentContext, tenant_scope
from app.db import SessionLocal
from app.models import (
    AnalysisFinding,
    AnalysisRun,
    AnalysisSchedule,
    AuditLog,
    EpssScore,
    IAMUser,
    Projects,
    ReportDelivery,
    ReportSubscription,
    SBOMComponent,
    SBOMSource,
    Tenant,
    TenantUser,
)
from app.schemas_reports import ReportPreferences, preferences_for
from app.services.email_sender import EmailDeliveryResult, EmailDeliveryStatus
from app.services.report_composer import compose_report
from app.services.report_cycles import (
    cadence_window,
    create_delivery,
    prepare_run_cycle,
    record_run_completion,
)
from app.services.report_rendering import render_email, render_pdf, render_xlsx, spreadsheet_value
from app.services.report_storage import configuration_errors
from app.settings import get_settings
from openpyxl import load_workbook
from pydantic import ValidationError
from sqlalchemy import func, select


@pytest.fixture
def fixture(client, monkeypatch, tmp_path):
    from app.routers.report_notifications import reader
    from app.workers import report_notifications as worker

    now = datetime.now(UTC)
    old = (now - timedelta(days=10)).isoformat()
    settings = get_settings()
    monkeypatch.setattr(settings, "tenant_role_assignment_mode", "LEGACY")
    db = SessionLocal()
    user = IAMUser(
        external_iam_user_id="report-user",
        email="verified@example.test",
        display_name="Verified",
        status="ACTIVE",
        email_verified=True,
        verification_required=False,
        email_verified_at=now,
        created_at=now,
        updated_at=now,
    )
    db.add(user)
    db.flush()
    member = TenantUser(
        tenant_id=1, user_id=user.id, role="TENANT_ADMIN", status="ACTIVE", created_at=now, updated_at=now
    )
    project = Projects(tenant_id=1, project_name="Reports", project_status=1, created_on=old)
    db.add_all([member, project])
    db.flush()
    root = SBOMSource(tenant_id=1, projectid=project.id, sbom_name="Original", sbom_version="1", created_on=old)
    db.add(root)
    db.flush()
    sbom = SBOMSource(
        tenant_id=1,
        projectid=project.id,
        sbom_name="<script>bad</script>=SUM(A1)",
        sbom_version="2",
        parent_id=root.id,
        created_on=old,
    )
    db.add(sbom)
    db.flush()
    component = SBOMComponent(
        tenant_id=1, sbom_id=sbom.id, name="library", version="2", is_duplicate=False, lifecycle_status="eol"
    )
    db.add(component)
    db.flush()
    runs = []
    for index, source in enumerate([root, sbom, sbom]):
        when = (now - timedelta(days=9 - index)).isoformat()
        run = AnalysisRun(
            tenant_id=1,
            sbom_id=source.id,
            project_id=project.id,
            run_status="FINDINGS",
            started_on=when,
            completed_on=when,
            total_components=1,
            total_findings=1,
        )
        db.add(run)
        db.flush()
        db.add(
            AnalysisFinding(
                tenant_id=1,
                analysis_run_id=run.id,
                vuln_id="CVE-2025-10001",
                component_name="library",
                component_id=component.id if source.id == sbom.id else None,
                component_version=source.sbom_version,
                severity="HIGH",
                score=9.0,
                fixed_versions='["3"]',
            )
        )
        runs.append(run)
    db.add(EpssScore(cve_id="CVE-2025-10001", epss=0.6, refreshed_at=now.isoformat()))
    db.commit()
    permissions = frozenset({"sbom:read", "analysis:read", "project:read", "product:read", "product:manage_schedule"})
    context = CurrentContext(
        user.id,
        user.external_iam_user_id,
        user.email,
        user.display_name,
        1,
        None,
        frozenset({"TENANT_ADMIN"}),
        permissions,
    )
    client.app.dependency_overrides[reader] = lambda: context
    private = tmp_path / "artifacts"
    private.mkdir(mode=0o700)
    worker_settings = settings.model_copy(
        update={
            "auth_enabled": True,
            "report_notifications_enabled": True,
            "email_delivery_enabled": True,
            "report_notification_base_url": "https://app.example.test",
            "report_artifact_storage_path": str(private),
            "email_from_address": "reports@example.test",
        }
    )
    monkeypatch.setattr(worker, "get_settings", lambda: worker_settings)
    import app.services.report_storage as storage

    monkeypatch.setattr(storage, "get_settings", lambda: worker_settings)
    sent = []
    monkeypatch.setattr(
        worker,
        "get_email_sender",
        lambda: SimpleNamespace(
            send_email=lambda message: (sent.append(message), EmailDeliveryResult(EmailDeliveryStatus.SENT))[1]
        ),
    )
    value = SimpleNamespace(
        db=db,
        client=client,
        user=user,
        member=member,
        project=project,
        root=root,
        sbom=sbom,
        runs=runs,
        context=context,
        worker=worker,
        settings=worker_settings,
        sent=sent,
        now=now,
        old=old,
    )
    yield value
    client.app.dependency_overrides.pop(reader, None)
    db.close()


def subscribe(f, **overrides):
    p = ReportPreferences(scope="SBOM", sbom_id=f.sbom.id, parts=["A", "B", "C", "D"], **overrides)
    sub = ReportSubscription(**p.database_values(), tenant_id=1, iam_user_id=f.user.id, created_on=f.old)
    f.db.add(sub)
    f.db.commit()
    return sub


def queued(f, sub):
    row = create_delivery(f.db, sub, start=f.old, end=f.now.isoformat(), sbom_ids=[f.sbom.id])
    f.db.commit()
    return row.id


@pytest.mark.parametrize(
    "values",
    [
        {"scope": "TENANT", "sbom_id": 1},
        {"scope": "SBOM"},
        {"scope": "PROJECT", "project_id": 1, "parts": ["B"]},
        {"scope": "TENANT", "parts": ["A"], "suppress_when_unchanged": True},
        {"scope": "TENANT", "timezone": "Not/AZone"},
        {"scope": "TENANT", "recipient_email": "attacker@example.test"},
        {"scope": "TENANT", "formats": ["HTML"]},
    ],
)
def test_preferences_reject_invalid(values):
    with pytest.raises(ValidationError):
        ReportPreferences(**values)


def test_all_parts_and_real_artifact_renderers(fixture):
    f = fixture
    sub = subscribe(f)
    with tenant_scope(f.context):
        report = compose_report(f.db, preferences_for(sub), tenant_id=1, cycle_start=f.old, cycle_end=f.now.isoformat())
    assert report["summary"]["total_findings"] == 1
    assert report["summary"]["severity"]["HIGH"] == 1
    assert report["summary"]["epss_outlook"]["coverage"] == 1
    item = report["sboms"][0]
    assert item["comparisons"]["B"]["run_a"]["id"] == f.runs[1].id
    assert item["comparisons"]["C"]["persistent_findings"][0]["age_days"] > 0
    assert item["comparisons"]["D"]["relationship"]["classification"] == "SAME_LINEAGE_DIFFERENT_VERSION"
    assert item["comparisons"]["D"]["relationship"]["version_a"] == "1"
    assert item["comparisons"]["D"]["newly_kev"]["status"] == "unavailable"
    _, plain, html = render_email(report)
    assert "<script>" not in html and "&lt;script&gt;" in html
    assert "Convention A" in plain and "coverage" in plain
    assert render_pdf(report).startswith(b"%PDF")
    book = load_workbook(BytesIO(render_xlsx(report)), read_only=True)
    assert book.sheetnames == ["Rollup", "Part A", "Part B", "Part C", "Part D", "Metadata"]
    assert all(cell.data_type != "f" for sheet in book for row in sheet for cell in row)
    book.close()
    assert spreadsheet_value(' =HYPERLINK("evil")').startswith("'")


def test_preview_crud_unique_and_no_delivery(fixture):
    f = fixture
    payload = ReportPreferences(scope="SBOM", sbom_id=f.sbom.id).model_dump()
    response = f.client.post("/api/report-subscriptions/preview", json=payload)
    assert response.status_code == 200, response.text
    assert "html_body" in response.json()
    assert f.db.scalar(select(func.count(ReportDelivery.id))) == 0
    response = f.client.post("/api/report-subscriptions", json=payload)
    assert response.status_code == 201, response.text
    identifier = response.json()["id"]
    assert f.client.post("/api/report-subscriptions", json=payload).status_code == 409
    assert f.client.patch(f"/api/report-subscriptions/{identifier}", json={"enabled": False}).status_code == 200
    assert f.client.post(f"/api/report-subscriptions/{identifier}/send-now").status_code == 409
    assert f.client.delete(f"/api/report-subscriptions/{identifier}").status_code == 204
    assert f.client.post("/api/report-subscriptions", json=payload).status_code == 201
    assert f.db.scalar(select(func.count(AuditLog.id)).where(AuditLog.action == "report.subscription.create")) == 2
    assert not f.sent


def test_tenant_scope_restricted_and_foreign_target_denied(fixture):
    f = fixture
    f.member.role = "VIEWER"
    f.db.commit()
    response = f.client.post("/api/report-subscriptions", json={"scope": "TENANT"})
    assert response.status_code == 403, response.text
    other = Tenant(id=2, name="Other", slug="other", status="ACTIVE", created_at=f.now, updated_at=f.now)
    f.db.add(other)
    f.db.flush()
    with tenant_scope(CurrentContext(0, "system", None, None, 2, None, frozenset(), frozenset())):
        sbom = SBOMSource(tenant_id=2, sbom_name="Secret")
        f.db.add(sbom)
        f.db.commit()
    response = f.client.post("/api/report-subscriptions", json={"scope": "SBOM", "sbom_id": sbom.id})
    assert response.status_code == 404


def test_successful_delivery_once_and_authenticated_download(fixture):
    f = fixture
    sub = subscribe(f)
    identifier = queued(f, sub)
    assert f.worker.generate(identifier, 1)["status"] == "SENT"
    assert f.worker.generate(identifier, 1)["status"] == "NOT_CLAIMED"
    assert len(f.sent) == 1
    assert str(f.sent[0]["To"]) == "verified@example.test"
    assert len(list(f.sent[0].iter_attachments())) == 2
    response = f.client.get("/api/report-deliveries")
    assert response.status_code == 200
    delivery = response.json()[0]
    assert len(delivery["artifacts"]) == 3
    assert "storage_path" not in response.text and "payload" not in response.text
    artifact = delivery["artifacts"][0]
    response = f.client.get(f"/api/report-deliveries/{identifier}/artifacts/{artifact['id']}")
    assert response.status_code == 200, response.text
    assert response.content.startswith(b"%PDF")
    f.db.expire_all()
    assert f.db.get(ReportSubscription, sub.id).last_delivered_at == f.now.isoformat()


@pytest.mark.parametrize(
    "change,expected",
    [("member", "MEMBERSHIP_REVOKED"), ("user", "RECIPIENT_NOT_VERIFIED"), ("paused", "SUBSCRIPTION_DISABLED")],
)
def test_send_time_rechecks(fixture, change, expected):
    f = fixture
    sub = subscribe(f)
    identifier = queued(f, sub)
    if change == "member":
        f.member.status = "DISABLED"
    elif change == "user":
        f.user.email_verified = False
    else:
        sub.enabled = False
    f.db.commit()
    result = f.worker.generate(identifier, 1)
    assert result == {"status": "SUPPRESSED", "error_code": expected}
    assert not f.sent


def test_disabled_local_mode_skips(fixture):
    f = fixture
    identifier = queued(f, subscribe(f))
    f.settings.auth_enabled = False
    assert f.worker.generate(identifier, 1)["status"] == "SKIPPED"
    assert not f.sent


def test_tenant_quota_and_size_links(fixture):
    f = fixture
    identifier = queued(f, subscribe(f))
    f.settings.report_max_attachment_bytes = 1
    assert f.worker.generate(identifier, 1)["status"] == "SENT"
    assert not list(f.sent[0].iter_attachments())
    body = f.sent[0].get_body(preferencelist=("plain",)).get_content()
    assert (
        "exceeds the attachment limit" in body and "https://app.example.test/settings/notifications?delivery=" in body
    )
    assert "storage_path" not in body


def test_controlled_retry_and_permanent_smtp_error(fixture, monkeypatch):
    f = fixture
    identifier = queued(f, subscribe(f))
    monkeypatch.setattr(
        f.worker,
        "get_email_sender",
        lambda: SimpleNamespace(
            send_email=lambda message: EmailDeliveryResult(EmailDeliveryStatus.FAILED, "SMTP_UNAVAILABLE")
        ),
    )
    result = f.worker.generate(identifier, 1)
    assert result["status"] == "PENDING"
    f.db.expire_all()
    row = f.db.get(ReportDelivery, identifier)
    assert row.attempt_count == 1 and row.next_attempt_at and not row.dispatch_started_at
    row.next_attempt_at = None
    f.db.commit()
    monkeypatch.setattr(
        f.worker,
        "get_email_sender",
        lambda: SimpleNamespace(
            send_email=lambda message: EmailDeliveryResult(EmailDeliveryStatus.FAILED, "SMTP_AUTHENTICATION_FAILED")
        ),
    )
    assert f.worker.generate(identifier, 1)["status"] == "FAILED"
    f.db.expire_all()
    assert f.db.get(ReportDelivery, identifier).attempt_count == 2


def test_overlapping_cycle_and_completion_barrier(fixture):
    f = fixture
    narrow = subscribe(f, cadence="ON_EVERY_RUN")
    broad = ReportSubscription(
        **ReportPreferences(scope="PROJECT", project_id=f.project.id, cadence="ON_EVERY_RUN").database_values(),
        tenant_id=1,
        iam_user_id=f.user.id,
        created_on=f.old,
    )
    f.db.add(broad)
    f.db.commit()
    prepare_run_cycle(f.db, [SimpleNamespace(sbom_id=f.sbom.id)], f.now.isoformat())
    rows = list(f.db.scalars(select(ReportDelivery).order_by(ReportDelivery.id)))
    selected = next(row for row in rows if row.subscription_id == narrow.id)
    assert selected.payload["expected"] == [f.sbom.id]
    assert next(row for row in rows if row.subscription_id == broad.id).status == "SKIPPED"
    assert f.worker.generate(selected.id, 1)["status"] == "WAITING_FOR_RUNS"
    record_run_completion(
        f.db, tenant_id=1, sbom_id=f.sbom.id, cycle=f.now.isoformat(), status="FINDINGS", run_id=f.runs[-1].id
    )
    assert f.worker.generate(selected.id, 1)["status"] == "SENT"
    assert len(f.sent) == 1


def test_concurrent_workers_do_not_double_send(fixture):
    f = fixture
    identifier = queued(f, subscribe(f))
    with ThreadPoolExecutor(max_workers=2) as pool:
        result = list(pool.map(lambda _: f.worker.generate(identifier, 1), range(2)))
    assert sorted(row["status"] for row in result) == ["NOT_CLAIMED", "SENT"]
    assert len(f.sent) == 1


def test_cadence_dst_month_boundary_and_missing_history(fixture):
    f = fixture
    sub = subscribe(f, cadence="DAILY", timezone="America/New_York")
    sub.created_on = "2026-01-01T00:00:00+00:00"
    start, end = cadence_window(sub, datetime(2026, 3, 9, 8, tzinfo=UTC))
    assert (datetime.fromisoformat(end) - datetime.fromisoformat(start)).total_seconds() == 23 * 3600
    sub.cadence = "MONTHLY"
    assert cadence_window(sub, datetime(2026, 4, 1, 12, tzinfo=UTC))[1].startswith("2026-04-01")
    p = ReportPreferences(scope="SBOM", sbom_id=f.root.id, parts=["A", "B", "C", "D"])
    report = compose_report(f.db, p, tenant_id=1, cycle_start=f.old, cycle_end=f.now.isoformat())
    assert all(value["status"] == "insufficient_history" for value in report["sboms"][0]["comparisons"].values())


def test_severity_floor_preserves_headline_and_scope_cap(fixture, monkeypatch):
    f = fixture
    report = compose_report(
        f.db,
        ReportPreferences(scope="SBOM", sbom_id=f.sbom.id, severity_floor="CRITICAL"),
        tenant_id=1,
        cycle_start=f.old,
        cycle_end=f.now.isoformat(),
    )
    assert report["summary"]["total_findings"] == 1 and report["sboms"][0]["A"]["findings"] == []
    broad = compose_report(
        f.db,
        ReportPreferences(scope="PROJECT", project_id=f.project.id),
        tenant_id=1,
        cycle_start=f.old,
        cycle_end=f.now.isoformat(),
    )
    assert broad["total_sboms"] == 1  # root is baseline, never double-counted in latest state


def test_tenant_schedule_cascade_and_paused_override(fixture):
    from app.services.schedule_resolver import find_due_targets, resolve_for_sbom

    f = fixture
    parent = AnalysisSchedule(tenant_id=1, scope="TENANT", cadence="DAILY", hour_utc=2, enabled=True, next_run_at=f.old)
    f.db.add(parent)
    f.db.commit()
    assert resolve_for_sbom(f.db, f.sbom.id).id == parent.id
    child = AnalysisSchedule(tenant_id=1, scope="SBOM", sbom_id=f.sbom.id, cadence="DAILY", hour_utc=2, enabled=False)
    f.db.add(child)
    f.db.commit()
    assert f.sbom.id not in {r.sbom_id for r in find_due_targets(f.db, f.now.isoformat())}


def test_configuration_rejects_unsafe_urls_and_paths(fixture):
    f = fixture
    assert configuration_errors(f.settings) == []
    for url in [
        "http://remote.example.test",
        "https://user:secret@app.example.test",
        "https://app.example.test?token=secret",
    ]:
        assert "REPORT_BASE_URL_REQUIRED" in configuration_errors(
            f.settings.model_copy(update={"report_notification_base_url": url})
        )
    assert "REPORT_PRIVATE_STORAGE_REQUIRED" in configuration_errors(
        f.settings.model_copy(update={"report_artifact_storage_path": "/"})
    )


def test_partial_comparison_failure_still_sends(fixture, monkeypatch):
    from app.services.compare_service import CompareService

    f = fixture
    original = CompareService.compare

    def fail_root(self, a, b):
        if a == f.runs[0].id:
            raise RuntimeError("sensitive internal payload must not be returned")
        return original(self, a, b)

    monkeypatch.setattr(CompareService, "compare", fail_root)
    identifier = queued(f, subscribe(f))
    assert f.worker.generate(identifier, 1)["status"] == "SENT"
    assert "sensitive internal" not in f.sent[0].as_string()
    assert "unavailable" in f.sent[0].get_body(preferencelist=("plain",)).get_content()


def test_quota_is_durable_between_sessions(fixture):
    f = fixture
    sub = subscribe(f)
    f.settings.report_max_emails_per_tenant_per_hour = 1
    identifier = queued(f, sub)
    assert f.worker.generate(identifier, 1)["status"] == "SENT"
    f.db.expire_all()
    second = create_delivery(
        f.db, sub, start=f.old, end=(f.now + timedelta(seconds=1)).isoformat(), sbom_ids=[f.sbom.id]
    )
    f.db.commit()
    assert f.worker.generate(second.id, 1)["status"] == "DEFERRED"
    assert len(f.sent) == 1


def test_permission_revoked_during_render_prevents_send(fixture, monkeypatch):
    f = fixture
    render = f.worker.render_attachments

    def revoke(report, formats):
        result = render(report, formats)
        with SessionLocal() as db:
            member = db.get(TenantUser, f.member.id)
            member.status = "DISABLED"
            db.commit()
        return result

    monkeypatch.setattr(f.worker, "render_attachments", revoke)
    identifier = queued(f, subscribe(f))
    assert f.worker.generate(identifier, 1)["status"] == "SUPPRESSED"
    assert not f.sent


def test_unknown_smtp_outcome_is_not_retried(fixture):
    f = fixture
    identifier = queued(f, subscribe(f))
    row = f.db.get(ReportDelivery, identifier)
    row.claimed_at = f.old
    row.dispatch_started_at = f.old
    row.attempt_count = 1
    f.db.commit()
    f.worker.dispatch_pending()
    f.db.expire_all()
    assert row.status == "FAILED" and row.error_code == "SMTP_OUTCOME_UNKNOWN"
    assert f.worker.generate(identifier, 1)["status"] == "NOT_CLAIMED"
    assert not f.sent


def test_preview_is_rate_limited(fixture):
    f = fixture
    f.db.add_all(
        [
            AuditLog(
                tenant_id=1,
                user_ref_id=f.user.id,
                action="report.preview",
                target_kind="report_subscription",
                created_at=f.now.isoformat(),
            )
            for _ in range(10)
        ]
    )
    f.db.commit()
    response = f.client.post("/api/report-subscriptions/preview", json={"scope": "SBOM", "sbom_id": f.sbom.id})
    assert response.status_code == 429


def test_report_routes_use_read_permission_not_tenant_settings_permission():
    from starlette.requests import Request
    from app.core.security import permission_for_request
    for method, path in [("POST", "/api/report-subscriptions"), ("PATCH", "/api/report-subscriptions/1"),
                         ("POST", "/api/report-subscriptions/preview"), ("GET", "/api/report-deliveries/1/artifacts/1")]:
        assert permission_for_request(Request({"type": "http", "method": method, "path": path, "headers": []})) == "sbom:read"


def test_read_only_member_can_subscribe_to_project(fixture):
    from app.core.security import get_current_tenant_context
    from dataclasses import replace
    f = fixture
    f.member.role = "VIEWER"
    f.db.commit()
    context = replace(f.context, roles=frozenset({"VIEWER"}))
    f.client.app.dependency_overrides[get_current_tenant_context] = lambda: context
    try:
        response = f.client.post("/api/report-subscriptions", json={"scope": "PROJECT", "project_id": f.project.id})
        assert response.status_code == 201, response.text
    finally:
        f.client.app.dependency_overrides.pop(get_current_tenant_context, None)


@pytest.mark.parametrize("sbom_count", [50, 250])
def test_report_reference_load(fixture, sbom_count):
    """Reference: 50 SBOM/4,901 findings; fan-out: 250 SBOM/24,901 findings."""
    import resource
    import sys
    import time

    from sqlalchemy import insert

    f = fixture
    ids = list(
        f.db.scalars(
            insert(SBOMSource).returning(SBOMSource.id),
            [
                {
                    "tenant_id": 1,
                    "projectid": f.project.id,
                    "sbom_name": f"Load {index}",
                    "created_on": f.old,
                    "is_active": True,
                }
                for index in range(sbom_count - 1)
            ],
        )
    )
    run_ids = list(
        f.db.scalars(
            insert(AnalysisRun).returning(AnalysisRun.id),
            [
                {
                    "tenant_id": 1,
                    "sbom_id": sid,
                    "project_id": f.project.id,
                    "run_status": "FINDINGS",
                    "started_on": (f.now - timedelta(days=day)).isoformat(),
                    "completed_on": (f.now - timedelta(days=day)).isoformat(),
                    "is_active": True,
                }
                for sid in ids
                for day in (2, 1)
            ],
        )
    )
    for start in range(0, len(run_ids), 50):
        f.db.execute(
            insert(AnalysisFinding),
            [
                {
                    "tenant_id": 1,
                    "analysis_run_id": rid,
                    "vuln_id": f"CVE-2025-{10000 + i}",
                    "component_name": f"dependency-{i}",
                    "component_version": "1",
                    "severity": "HIGH",
                    "is_active": True,
                }
                for rid in run_ids[start : start + 50]
                for i in range(100)
            ],
        )
    f.db.commit()
    start = time.monotonic()
    with tenant_scope(f.context):
        report = compose_report(
            f.db,
            ReportPreferences(scope="PROJECT", project_id=f.project.id, parts=["A", "B", "C", "D"]),
            tenant_id=1,
            cycle_start=f.old,
            cycle_end=f.now.isoformat(),
        )
    f.db.close()
    pdf, xlsx = render_pdf(report), render_xlsx(report)
    elapsed = time.monotonic() - start
    memory_mb = resource.getrusage(resource.RUSAGE_SELF).ru_maxrss / (1024**2 if sys.platform == "darwin" else 1024)
    print(
        f"REPORT_LOAD sboms={sbom_count} findings={report['summary']['total_findings']} elapsed_seconds={elapsed:.2f} peak_mb={memory_mb:.1f} pdf_bytes={len(pdf)} xlsx_bytes={len(xlsx)}"
    )
    assert report["included_sboms"] == sbom_count
    assert elapsed < (120 if sbom_count == 50 else 300)
    assert memory_mb < 512
