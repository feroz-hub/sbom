import importlib.util,json,hashlib,io,csv,zipfile
from pathlib import Path
from urllib.parse import urlencode
spec=importlib.util.spec_from_file_location('m',Path(__file__).with_name('live-common.py'));m=importlib.util.module_from_spec(spec);spec.loader.exec_module(m);f=json.loads((m.OUT/'round3-scope.json').read_text());pid=f['project_id'];product=f['product_id'];sid,inactive=f['sbom_ids'];file='scope-exports.json'
def rec(id,steps,expected,r,ok,role='admin'):m.rec(id,steps,expected,m.payload(r) if hasattr(r,'status_code') else r,ok,role,file)
assert m.call('/api/projects/'+pid).json()['project_name']==f['name']
r=m.call('/api/sboms/'+sid+'/analyze/stream','POST',{'sources':['OSV']});rec('R3-OWNED-ANALYSIS','Analyze fresh own scope','200 complete',{'status':r.status_code,'complete':'event: complete' in r.text},r.status_code==200 and 'event: complete' in r.text)
r=m.call('/api/sboms/'+sid+'/analysis-runs');assert r.status_code==200;runs=r.json();run=next(x for x in runs if x['run_status'] in ['OK','FINDINGS','PARTIAL']);runid=run['id'];f['analysis_run_id']=runid;(m.OUT/'round3-scope.json').write_text(json.dumps(f,indent=2))
scope={'project_id':pid,'product_id':product};query=urlencode(scope)
before=m.call('/dashboard/summary?'+query);rec('R3-DASHBOARD-OWNED-BEFORE','Read fresh two-SBOM scoped posture','200 two stored active SBOMs',before,before.status_code==200 and before.json()['posture']['total_sboms']==2)
r=m.call('/api/sboms/'+inactive+'/lifecycle','POST',{'status':'INACTIVE','reason':'Round 3 task-owned inactive exclusion fixture'});rec('R3-DASHBOARD-INACTIVATE-OWNED','Deactivate only second freshly created SBOM','200 INACTIVE',r,r.status_code==200 and r.json()['lifecycle_status']=='INACTIVE')
posture=None
for role in ['admin','analyst','developer','viewer']:
 r=m.call('/dashboard/summary?'+query,role=role);body=r.json().get('posture',{});posture=body if posture is None else posture;rec('R3-DASHBOARD-INACTIVE-EXCLUDED-'+role,'Read same project/application scope after owned deactivation','200 one eligible stored SBOM and identical scoped metrics',r,r.status_code==200 and body.get('total_sboms')==1 and body==posture,role)
r=m.call('/dashboard/filter-options/sboms?'+query);items=r.json();items=items if isinstance(items,list) else items.get('items',[]);rec('R3-DASHBOARD-INACTIVE-OPTIONS','Read actual dependent SBOM choices','Active UUID present, inactive UUID excluded',r,r.status_code==200 and sid in [x['id'] for x in items] and inactive not in [x['id'] for x in items])
r=m.call('/dashboard/summary?'+query,role='foreign');rec('R3-DASHBOARD-OWNED-FOREIGN','Foreign user supplies known R3 UUID hierarchy','404',r,r.status_code==404,'foreign')
exports=[('pdf','/api/pdf-report','POST',{'runId':runid,'title':'Round 3 synthetic UAT report','filename':'round3.pdf'}),('sarif','/api/analysis-runs/'+str(runid)+'/export/sarif','GET',None),('csv','/api/analysis-runs/'+str(runid)+'/export/csv','GET',None),('xlsx','/api/sboms/'+sid+'/reports/vulnerabilities.xlsx','GET',None)]
for fmt,path,method,body in exports:
 r=m.call(path,method,body);ok=r.status_code==200;details={'status':r.status_code,'bytes':len(r.content),'contentType':r.headers.get('content-type'),'download':r.headers.get('content-disposition')}
 if ok:
  if fmt=='pdf':ok=r.content.startswith(b'%PDF');details['headerValid']=ok
  elif fmt=='sarif':data=r.json();details.update(version=data.get('version'),results=len(data['runs'][0]['results']));ok=data['version']=='2.1.0' and details['results']>0
  elif fmt=='csv':data=list(csv.DictReader(io.StringIO(r.text)));details.update(rows=len(data),fields=list(data[0]) if data else []);ok=bool(data) and {'vuln_id','severity','component_name'}.issubset(data[0])
  else:
   from openpyxl import load_workbook
   book=load_workbook(io.BytesIO(r.content),read_only=True);details['sheets']=book.sheetnames;details['sheetRows']={s.title:s.max_row for s in book};ok=bool(book.sheetnames);book.close()
  (m.OUT/('report-'+fmt+'.'+fmt)).write_bytes(r.content);details['sha256']=hashlib.sha256(r.content).hexdigest()
 else:details['error']=m.payload(r)
 rec('R3-EXPORT-'+fmt,'Generate and inspect real '+fmt+' artifact','200 valid content and downloadable headers',details,ok)
 for role in ['viewer','foreign']:
  denied=m.call(path,method,body,role=role);expected=403 if role=='viewer' else 404;rec('R3-EXPORT-'+fmt+'-'+role,'Direct '+role+' export of known owned resource',str(expected),denied,denied.status_code==expected,role)
# FDA requires both completed finding and lifecycle evidence; record actual prerequisites.
r=m.call('/api/sboms/'+sid+'/lifecycle/refresh','POST');rec('R3-FDA-LIFECYCLE-REFRESH','Request supported lifecycle evidence for own fresh fixture','Record actual configured lifecycle availability',r,r.status_code==200)
body={'selections':[{'sbom_id':sid,'findings_analysis_run_id':runid}],'metadata':{'device_name':'Synthetic UAT Device','manufacturer_sponsor':'Round3 Local Test','device_software_version':'1.0','author_of_sbom_data':'UAT','prepared_by':'UAT'}}
r=m.call('/api/projects/'+pid+'/reports/fda-510k-sbom/export','POST',body);ok=r.status_code==200 and r.content.startswith(b'PK');details=m.payload(r)
if ok:
 (m.OUT/'report-fda.xlsx').write_bytes(r.content);from openpyxl import load_workbook
 b=load_workbook(io.BytesIO(r.content),read_only=True);details={'status':200,'bytes':len(r.content),'sheets':b.sheetnames,'rows':{s.title:s.max_row for s in b}};b.close()
rec('R3-EXPORT-FDA','Export required FDA workbook with synthetic metadata and explicit finding run','200 populated workbook; otherwise preserve actual evidence prerequisite',details,ok)
